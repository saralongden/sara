"""Build an editable A4 PPTX of the women's pant size chart for Canva import.

Every label, number, rule and fill is its own native text box or shape, so the
whole chart stays editable once imported into Canva.

Usage: python3 tools/build_size_chart_pptx.py <boxes.json> <out.pptx>
boxes.json is the layout measured from size-charts/womens-pant-size-chart.html.
"""
import json
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Mm, Pt

PX = 9525  # EMU per CSS pixel (96 dpi)
NAVY, GOLD, PEARL, TINT, PALE_GOLD = "1F2A44", "9E7A38", "F8F5F1", "F1EBE2", "E4D3B0"
LINE = "DCCDB0"  # gold at 35% on pearl
SERIF, SANS = "Playfair Display", "Montserrat"

d = json.load(open(sys.argv[1]))
prs = Presentation()
prs.slide_width, prs.slide_height = Mm(210), Mm(297)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def px(v):
    return Emu(int(round(v * PX)))


def rect(x, y, w, h, fill=None, line=None, line_w=0.75):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    s.shadow.inherit = False
    style = s._element.find(qn("p:style"))
    if style is not None:  # drop theme style so Canva/PowerPoint add no shadow or outline
        s._element.remove(style)
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = RGBColor.from_string(line); s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    return s


def text(b, t, font, size_pt, color, bold=False, italic=False, spacing_em=0, align=PP_ALIGN.CENTER, pad=0):
    tb = slide.shapes.add_textbox(px(b["x"] - pad), px(b["y"]), px(b["w"] + 2 * pad), px(b["h"]))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        f = r.font; f.name = font; f.size = Pt(size_pt); f.bold, f.italic = bold, italic
        f.color.rgb = RGBColor.from_string(color)
        if spacing_em:
            r._r.get_or_add_rPr().set("spc", str(int(round(spacing_em * size_pt * 100))))
    return tb


W, H = d["page"]["w"], d["page"]["h"]
bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor.from_string(PEARL)
m = 30.2  # 8mm frame inset
rect(m, m, W - 2 * m, H - 2 * m, line=LINE, line_w=0.75)

# Logo lockup
text(d["small"], "GOLF FASHION", SANS, 8.5, GOLD, spacing_em=0.42, align=PP_ALIGN.LEFT, pad=0)
text(d["women"], "WOMEN", SERIF, 44, NAVY, spacing_em=0.02, align=PP_ALIGN.LEFT)
row, on = d["onrow"], d["on"]
cy = row["y"] + row["h"] / 2
gap = 18.9  # 5mm between rule and "on"
rect(row["x"], cy, on["x"] - gap - row["x"], 0.75, fill=GOLD)
rect(on["x"] + on["w"] + gap, cy, row["x"] + row["w"] - (on["x"] + on["w"] + gap), 0.75, fill=GOLD)
text(on, "on", SERIF, 30, GOLD, italic=True, pad=10)
text(d["course"], "COURSE", SERIF, 44, NAVY, spacing_em=0.02, align=PP_ALIGN.RIGHT)
text(d["since"], "SINCE 2004", SANS, 8.5, GOLD, spacing_em=0.42, align=PP_ALIGN.RIGHT)

# Title
text(d["h1"], "WOMEN’S PANT SIZING", SANS, 17, NAVY, spacing_em=0.38, pad=20)
text(d["sub"], "SIZE CONVERSION GUIDE", SANS, 8.5, GOLD, spacing_em=0.4, pad=20)

# Group labels
for g in d["group"]:
    if not g["t"]:
        continue
    text({**g, "h": g["h"] - 6}, g["t"], SANS, 7.5, GOLD, spacing_em=0.38)

# Header row (skip the spacer column)
hd = d["head"]
for seg in (hd[:4], hd[5:]):
    rect(seg[0]["x"], seg[0]["y"], seg[-1]["x"] + seg[-1]["w"] - seg[0]["x"], seg[0]["h"], fill=NAVY)
for i, c in enumerate(hd):
    if not c["t"]:
        continue
    brand = i > 4
    text(c, c["t"], SANS, 7.5 if brand else 8.5, PEARL if brand else PALE_GOLD,
         spacing_em=0.14 if brand else 0.24)

# Body rows: banding, rules, then values
for ri, cells in enumerate(d["rows"]):
    fill = TINT if ri % 2 else None
    for seg in (cells[:4], cells[5:]):
        x0, x1 = seg[0]["x"], seg[-1]["x"] + seg[-1]["w"]
        c = seg[0]
        if fill:
            rect(x0, c["y"], x1 - x0, c["h"], fill=fill)
        rect(x0, c["y"] + c["h"] - 0.75, x1 - x0, 0.75, fill=LINE)
    for ci, c in enumerate(cells):
        if ci == 4:
            continue
        if c["t"] and c["t"] != "–":
            text(c, c["t"], SERIF, 15, NAVY, bold=(ci == 0))
first, last = d["rows"][0][5], d["rows"][-1][5]
rect(first["x"], first["y"], 0.75, last["y"] + last["h"] - first["y"], fill=LINE)

# Footer
text(d["note"], "Sizes are a guide only — fit varies by style.\nOur team is happy to help you find your perfect fit.",
     SANS, 8, NAVY, spacing_em=0.06, pad=20)
text(d["foott"], "WOMEN ON COURSE · GOLF FASHION", SANS, 7.5, GOLD, spacing_em=0.4, pad=20)

prs.save(sys.argv[2])
