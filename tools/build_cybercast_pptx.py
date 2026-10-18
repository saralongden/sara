# Editable A4 PPTX of the Cybercast sheet for Canva: text boxes, shapes and the QR code image.
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
PX = 3048  # EMU per pixel at 300 dpi (A4 = 2480 x 3508 px)
NAVY, GOLD, GOLD_B, PEARL = "1F2A44", "9E7A38", "D9B66A", "F9F6F1"
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(2480 * PX), Emu(3508 * PX)
s = prs.slides.add_slide(prs.slide_layouts[6])
px = lambda v: Emu(int(round(v * PX)))
def rect(x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, px(x), px(y), px(w), px(h)); r.shadow.inherit = False
    if fill: r.fill.solid(); r.fill.fore_color.rgb = RGBColor.from_string(fill)
    else: r.fill.background()
    if line: r.line.color.rgb = RGBColor.from_string(line); r.line.width = Pt(lw)
    else: r.line.fill.background()
    return r
def text(x, y, w, h, t, face, size_px, color, bold=False, spacing=0, align=PP_ALIGN.CENTER):
    tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h)); tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = t
    f = r.font; f.name = face; f.size = Pt(size_px * 0.24); f.bold = bold; f.color.rgb = RGBColor.from_string(color)
    if spacing: r._r.get_or_add_rPr().set("spc", str(int(spacing * 0.24 * 100)))
    return tb
rect(0, 0, 2480, 3508, fill=PEARL)
rect(70, 70, 2340, 3368, line=GOLD, lw=1.0); rect(100, 100, 2280, 3308, line=GOLD, lw=0.5)
cx = 1240
# Wordmark
text(cx - 520, 175, 600, 50, "GOLF FASHION", "Montserrat", 25, GOLD, spacing=9, align=PP_ALIGN.LEFT)
text(cx - 600, 200, 1200, 200, "WOMEN", "Playfair Display", 158, NAVY, bold=True)
rect(cx - 273, 401, 189, 4, fill=GOLD); rect(cx + 84, 401, 189, 4, fill=GOLD)
text(cx - 150, 330, 300, 150, "on", "Pinyon Script", 126, GOLD)
text(cx - 560, 420, 1200, 200, "COURSE", "Playfair Display", 158, NAVY, bold=True)
text(cx - 100, 595, 600, 45, "SINCE 2004", "Montserrat", 25, GOLD, spacing=9, align=PP_ALIGN.RIGHT)
# Navy band
rect(100, 690, 2280, 320, fill=NAVY)
text(100, 720, 2280, 70, "MEMBER GUEST DAY", "Montserrat", 48, GOLD_B, bold=True, spacing=16)
text(100, 775, 2280, 170, "SPIN THE WHEEL", "Playfair Display", 150, PEARL, bold=True)
text(100, 935, 2280, 50, "TUESDAY 20 OCTOBER 2026  &  WEDNESDAY 21 OCTOBER 2026", "Montserrat", 34, PEARL, spacing=8)
# Offer
text(150, 1050, 2180, 100, "Follow us, Scan the QR code and", "Playfair Display", 76, NAVY, bold=True)
text(150, 1145, 2180, 115, "WIN up to 40% off all Women's apparel", "Playfair Display", 92, GOLD, bold=True)
# QR
text(100, 1270, 2280, 70, "SCAN TO WIN", "Montserrat", 60, GOLD, bold=True, spacing=20)
rect(cx - 776, 1384, 1552, 1552, line=GOLD, lw=1.5)
s.shapes.add_picture("women-on-course-guest-day/canva/spin-qr-1480.png", px(cx - 740), px(1420), px(1480), px(1480))
# Steps
steps = [("1", "Follow us", "@womenoncoursegolf"), ("2", "Scan the QR code", "and spin the wheel"), ("3", "Show your discount", "at the pro shop till")]
colw = (2480 - 400) / 3
for i, (n, t, sub) in enumerate(steps):
    x = 200 + colw * i + colw / 2
    rect(x - 46, 3014, 92, 92, line=GOLD, lw=1.0, shape=MSO_SHAPE.OVAL)
    text(x - 46, 3014, 92, 92, n, "Playfair Display", 58, GOLD, bold=True)
    text(x - colw / 2, 3130, colw, 80, t, "Playfair Display", 60, NAVY, bold=True)
    text(x - colw / 2, 3210, colw, 60, sub, "Montserrat", 38, NAVY)
text(150, 3335, 2180, 60, "One spin per person. Discount on Women's apparel in the pro shop, for a single purchase on 20 or 21 October 2026, while stocks last.", "Montserrat", 30, NAVY)
prs.save("/home/user/sara/women-on-course-guest-day/canva/woc-spin-the-wheel-A4-editable.pptx")
print("saved")
