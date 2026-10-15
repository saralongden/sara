"""Build an editable PPTX of the flyer for Canva import (text boxes, shapes, PNGs).

Usage: python3 tools/build_pptx.py <boxes.json> <out.pptx>
boxes.json is the element layout measured from new-arrivals.html.
"""
import json
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 9525  # EMU per CSS pixel (96 dpi)
OXFORD, ANTIQUE, PAPER = "293349", "AF8A53", "FFFDFA"
BLUSH, CREAM, CARD_LINE = "EDE4DB", "F7F2EC", "E6D9C6"

d = json.load(open(sys.argv[1]))
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(d["poster"]["w"] * PX), Emu(d["poster"]["h"] * PX)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def px(v):
    return Emu(int(round(v * PX)))


def rect(x, y, w, h, fill=None, line=None, line_w=0.75):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    s.shadow.inherit = False
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = RGBColor.from_string(line); s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    return s


def text(x, y, w, h, runs, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP):
    """runs: list of paragraphs; each paragraph a list of (text, font, size_px, color, bold, italic, spacing_em)."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, font, size, color, bold, italic, spc in para:
            r = p.add_run(); r.text = t
            f = r.font; f.name = font; f.size = Pt(size * 0.75)
            f.bold, f.italic = bold, italic
            f.color.rgb = RGBColor.from_string(color)
            if spc:
                r._r.get_or_add_rPr().set("spc", str(int(spc * size * 0.75 * 100)))
    return tb


W, H = d["poster"]["w"], d["poster"]["h"]

# Background: one full-page shape fading from blush (top) to cream, no seams
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, px(W), px(H))
bg.shadow.inherit = False
bg.line.fill.background()
bg.fill.gradient(); bg.fill.gradient_angle = 270
stops = bg.fill.gradient_stops
stops[0].color.rgb = RGBColor.from_string(BLUSH); stops[0].position = 0.0
stops[1].color.rgb = RGBColor.from_string(CREAM); stops[1].position = 0.55

# Double gold frame
f = d["frame"]
rect(f["x"], f["y"], f["w"], f["h"], line="CDB896", line_w=0.75)
rect(f["x"] + 6, f["y"] + 6, f["w"] - 12, f["h"] - 12, line="E4D8C3", line_w=0.75)

# Logo
lg = d["logo"]
slide.shapes.add_picture("rmgc-logo-primary.png", px(lg["x"]), px(lg["y"]), px(lg["w"]), px(lg["h"]))

# IN STORE NOW with rules either side
e = d["eyebrow"]
cy = e["y"] + e["h"] / 2
rect(e["x"], cy, 70, 1, fill="C9B48E")
rect(e["x"] + e["w"] - 70, cy, 70, 1, fill="C9B48E")
text(0, e["y"] - 1, W, e["h"] + 4, [[("IN STORE NOW", "Montserrat", 13, OXFORD, True, False, 0.5)]])

# Headline block
t = d["title"]
text(0, t["y"], W, t["h"], [[("New Arrivals", "Libre Caslon Display", 148, ANTIQUE, False, False, 0)]])
s = d["subtitle"]
text(0, s["y"], W, s["h"], [[("The Women’s Collection", "Libre Caslon Text", 40, OXFORD, False, True, 0)]])
b = d["brand"]
text(0, b["y"], W, b["h"], [[("TRAVIS", "Montserrat", 38, OXFORD, True, False, 0.5),
                             ("MATHEW", "Montserrat", 38, OXFORD, False, False, 0.5)]])
se = d["season"]
text(0, se["y"], W, se["h"] + 4, [[("SPRING · SUMMER ’26", "Libre Caslon Text", 13, ANTIQUE, False, False, 0.45)]])

# Product cards and cut-out images
for c in d["cards"]:
    rect(c["x"], c["y"], c["w"], c["h"], fill=PAPER, line=CARD_LINE)
for im in d["imgs"]:
    slide.shapes.add_picture(im["src"], px(im["x"]), px(im["y"]), px(im["w"]), px(im["h"]))

# Gold collection tile
ft = d["feature"]
rect(ft["x"], ft["y"], ft["w"], ft["h"], fill=ANTIQUE)
rect(ft["x"] + 12, ft["y"] + 12, ft["w"] - 24, ft["h"] - 24, line="D7C5A5")
k, big = d["fk"], d["fbig"]
text(ft["x"], k["y"], ft["w"], k["h"] + 4, [[("TRAVISMATHEW", "Montserrat", 17, PAPER, False, False, 0.45)]])
text(ft["x"], big["y"], ft["w"], big["h"] + 10, [
    [("Discover", "Libre Caslon Text", 30, PAPER, False, True, 0)],
    [("the full collection in store", "Libre Caslon Display", 50, PAPER, False, False, 0)],
])

prs.save(sys.argv[2])
print("saved", sys.argv[2], len(slide.shapes), "elements")
